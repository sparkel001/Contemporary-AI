"""Build the editable Chinese report from independently rechecked raw trials."""
import json
from pathlib import Path
import re
import statistics
import sys

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from audit_results import audit, independent_summary

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'src'))
from search_lab.examples import MAZE, TERRAIN, TREASURE_MAP

LABEL = {'zero':'h=0', 'manhattan':'修正曼哈顿', 'topology':'拓扑下界',
         'minimax':'Minimax', 'left':'AB 原顺序', 'right':'AB 逆序',
         'ordered':'AB 排序', 'cache':'AB 缓存', 'ordered_cache':'AB 排序+缓存'}


def main():
    audit_result = audit()
    data, _ = independent_summary(ROOT/'results/benchmarks')
    config = json.loads((ROOT/'results/benchmarks/config.json').read_text(encoding='utf-8'))
    branches = json.loads((ROOT/'results/benchmarks/decision_branches.json').read_text(encoding='utf-8'))
    examples = json.loads((ROOT/'results/examples.json').read_text(encoding='utf-8'))
    test_log = (ROOT/'results/pytest.txt').read_text(encoding='utf-8-sig')
    tests = re.search(r'(\d+) passed in ([\d.]+)s', test_log)
    if not tests:
        raise ValueError('No successful pytest record')
    def r(case, mode):
        return data[case, mode]
    def us(row, field='core_ns'):
        return f"{row[field]/1000:.2f}"
    def interval(row):
        return f"{us(row)} / {row['core_ns_iqr']/1000:.2f}"
    doc = Document()
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.8)
    sec.left_margin, sec.right_margin = Cm(2.1), Cm(2.1)
    sec.header_distance, sec.footer_distance = Cm(.8), Cm(.8)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Times New Roman', Pt(10.5)
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True
    for name, size in [('Title',24), ('Subtitle',14), ('Heading 1',15), ('Heading 2',12), ('Heading 3',11)]:
        style = doc.styles[name]
        style.font.name, style.font.size = 'Times New Roman', Pt(size)
        style.font.color.rgb = RGBColor(0,0,0)
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '黑体')
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(6)
    doc.styles['Title'].paragraph_format.space_after = Pt(15)
    doc.styles['Caption'].font.size = Pt(9)
    doc.styles['Caption'].font.color.rgb = RGBColor(0,0,0)
    doc.styles['Caption']._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
    header = sec.header.paragraphs[0]
    header.text = '当代人工智能实验课  实验二 搜索算法'
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0,0,0)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('第 ')
    for instruction, tail in [('PAGE',' 页 / 共 '), ('NUMPAGES',' 页')]:
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), instruction)
        footer._p.append(field)
        footer.add_run(tail)
    doc.core_properties.title = '当代人工智能实验二 搜索算法实验报告'
    doc.core_properties.subject = 'A星搜索 极大极小搜索 Alpha-Beta剪枝 性能对照'
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''

    def p(text, bold=False):
        para = doc.add_paragraph()
        run = para.add_run(text)
        run.bold = bold
        para.paragraph_format.first_line_indent = Cm(.74)
        return para
    def heading(text, level=1):
        return doc.add_heading(text, level)
    def page():
        doc.add_page_break()
    def equation(text):
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        math = OxmlElement('m:oMath')
        run = OxmlElement('m:r')
        content = OxmlElement('m:t')
        content.text = text
        run.append(content)
        math.append(run)
        para._p.append(math)
    def caption(text):
        para = doc.add_paragraph(text, 'Caption')
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(4)
        para.paragraph_format.space_after = Pt(4)
        return para
    def table(headers, rows, widths=None, title=None, fontsize=9):
        if title:
            cap = caption(title)
            cap.paragraph_format.keep_with_next = True
        t = doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        widths = widths or [16.8/len(headers)]*len(headers)
        for col, width in zip(t.columns, widths):
            col.width = Cm(width)
        for i, text in enumerate(headers):
            t.rows[0].cells[i].text = str(text)
        for values in rows:
            cells = t.add_row().cells
            for cell, value in zip(cells, values):
                cell.text = str(value)
        borders = OxmlElement('w:tblBorders')
        for edge in ('top','left','bottom','right','insideH','insideV'):
            element=OxmlElement('w:'+edge)
            element.set(qn('w:val'),'single'); element.set(qn('w:sz'),'4'); element.set(qn('w:color'),'D9D9D9')
            borders.append(element)
        t._tbl.tblPr.append(borders)
        for idx,row in enumerate(t.rows):
            pr = row._tr.get_or_add_trPr()
            pr.append(OxmlElement('w:cantSplit'))
            if idx==0:
                pr.append(OxmlElement('w:tblHeader'))
            for i,cell in enumerate(row.cells):
                cell.width = Cm(widths[i])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                margins=OxmlElement('w:tcMar')
                for edge in ('top','bottom','left','right'):
                    el=OxmlElement('w:'+edge); el.set(qn('w:w'),'55'); el.set(qn('w:type'),'dxa'); margins.append(el)
                cell._tc.get_or_add_tcPr().append(margins)
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'DBE5F1' if idx==0 else 'FFFFFF' if idx%2 else 'F7F8FA')
                cell._tc.get_or_add_tcPr().append(shade)
                for para in cell.paragraphs:
                    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    para.paragraph_format.space_after=Pt(1)
                    para.paragraph_format.space_before=Pt(1)
                    para.paragraph_format.line_spacing=1.1
                    for run in para.runs:
                        run.font.size=Pt(fontsize);run.bold=idx==0
        return t
    def figure(name, title, width=16.2):
        para=doc.add_paragraph()
        para.alignment=WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next=True
        shape=para.add_run().add_picture(str(ROOT/'results/figures'/name), width=Cm(width))
        shape._inline.docPr.set('descr',title)
        caption(title)

    # Page 1: title and purpose, rather than an otherwise empty cover.
    title=doc.add_paragraph('当代人工智能实验课', 'Subtitle');title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    title=doc.add_paragraph('实验二 搜索算法实验报告','Title');title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for text in ('姓名：________________    学号：________________', '实验日期：2026年10月19日'):
        para=doc.add_paragraph(text);para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    heading('摘要',2)
    p('本实验实现四方向 A* 迷宫寻路、加权地形最小代价搜索、Minimax 与 Alpha-Beta 博弈树求值，以及四宝藏竞争决策。四题官方结果均得到复现，综合题在 H=2 时选择 T2、收益差为 −1，在 H=4 时选择 T3、收益差为 3。')
    p('性能实验覆盖 36 张地图、7 棵树和两个综合题搜索深度，共 148 个配置、2,220 条正式计时记录。拓扑下界、节点排序与状态缓存分别接受消融验证。结果表明，更强启发可减少扩展，但其预处理可能增加单次总成本；即时收益排序减少了官方综合题的访问数；完整状态缓存在该实例中命中率为 0%，没有产生搜索收益。')
    heading('1 实验目的与任务')
    p('实验依据课程任务书第 3—15 页的输入、输出和规则实现，第 16—17 页所强调的对照实验与原因分析作为评价依据。研究重点是保持最优性与博弈语义不变时，启发信息、访问顺序、预处理及缓存如何影响搜索工作量。')
    table(['任务','关键约束','官方输出'],[
        ['迷宫寻路','四方向、单位移动代价、完整路径','7 步最短路径'],
        ['加权地形','费用为进入格权值，进入 E 为 0','最小代价 7'],
        ['博弈树','根 MAX，原始接口从左到右遍历','根值 5'],
        ['宝藏竞争','共享地图、每步仅当前玩家移动','H=2：T2/−1；H=4：T3/3']],
        [2.6,8.1,6.1], '表1 四项任务与验收结果')
    p('程序将算法、实验运行和图表展示分离。官方算法接口保留原返回类型，统计信息通过独立数据结构记录；实验扩展不替换官方规则，也不使用预期输出参与搜索。')

    page();heading('2 实验环境 数据与评价指标')
    table(['项目','设置'],[
        ['算法环境',config['python'].split(' (')[0]+'；'+config['platform']],
        ['依赖','算法与实验：Python 标准库；pytest 8.4.2；matplotlib '+config['packages']['matplotlib']],
        ['随机种子','主种子 20261019；地图种子 20261019—20261021'],
        ['重复测量','每配置预热 2 次，正式重复 15 次；每轮打乱算法顺序'],
        ['计时方法','perf_counter_ns；中位数、Q1、Q3、IQR 及极值'],
        ['正确性基准','独立 O(V²) Dijkstra、BFS、Bellman-Ford、完整 Minimax']],
        [3.1,13.7], '表2 实验环境与控制条件')
    p('随机地图为 24×24，障碍生成概率为 0.10、0.25、0.40，加权场景的非障碍数字取 1—3 或 1—9。每个概率使用三个固定种子，权重变化时保持障碍布局一致；不强制连通。额外固定输入包括不可达、等价最短路、折返通道和零权地形。树实验使用分支数 3、最大深度 3/5/7，非平衡树以 0.22 的概率提前终止为叶。')
    table(['指标','精确定义'],[
        ['generated','A* 插入项数，含初始堆项；博弈为实际递归子分支次数'],
        ['visited / expanded','有效出队或递归进入次数 / 实际展开后继次数；目标和终止叶不展开'],
        ['heap_pop_calls','所有 heappop 次数 = visited + stale_pops'],
        ['heap_operations','实际 heappush 与 heappop 调用之和；初始单项列表不调用 heappush'],
        ['leaf_evaluations','实际搜索终止评价；额外排序样本叶另计 ordering_leaves'],
        ['cutoffs / skipped_nodes','确实省略兄弟的事件数 / 省略分支完整展开后的节点发生次数'],
        ['cache_hit_rate','命中次数 / 探测次数；命中可能只是收紧窗口'],
        ['prep / core / total','预处理时间 / 核心时间 / 每次二者之和，再分别求统计量']],
        [4,12.8], '表3 统计口径',fontsize=9)
    p('输入构建、正确性基准计算与写文件不计入核心时间。BFS 下界构建和树结构大小统计计入预处理；排序探测和缓存维护计入核心时间。距离矩阵在综合题各次搜索之前统一构建，耗时单独写入 config.json，而非计入每次博弈决策。未控制 CPU 频率或关闭 GC，微秒级时间仅作为辅助证据。')

    page();heading('3 A星算法原理与实现')
    heading('3.1 状态 代价与优先队列',2)
    p('网格状态为坐标 n=(row,col)。g(n) 表示当前已发现的从起点到 n 的最小费用，h(n) 为剩余费用下界，按 f(n) 最小的候选优先出队。起点 g=0；题一每条边费用为 1，题二从 p 到 q 的边费用为 q 的进入权值。')
    equation('f(n) = g(n) + h(n)')
    equation('g_new(q) = g(p) + c(p,q)')
    p('堆项保存 (f,入队序号,g,坐标)。同 f 值按 FIFO 排序，四方向顺序固定为右、下、左、上。当新 g 严格小于记录值时更新父节点并重新入堆；出队项的 g 已过期时丢弃。算法不依赖永久关闭集合，因此允许更优路径重新进入。严格改进条件避免零费用环导致反复更新。')
    table(['步骤','处理'],[
        ['初始化','distance[start]=0，起点入堆，parent 为空'],
        ['取候选','弹出最小 f 项；过期项仅增加 stale_pops'],
        ['到达目标','沿 parent 逆向回溯并反转，得到完整路径'],
        ['扩展','枚举四邻格，跳过越界与障碍，执行严格松弛'],
        ['失败','堆耗尽后题一返回 []，题二返回 −1']],
        [3,13.8],'表4 A* 的实现流程')
    heading('3.2 最优性条件与边界处理',2)
    p('可采纳性要求 h(n) 不超过到目标的真实最小费用，且目标处 h=0。一致性进一步要求每条边满足下式，使 f 沿路径不下降。本文使用的启发函数均满足该条件；在非负边权下，目标以最小 f 出队时可以返回最优结果。')
    equation('0 ≤ h(n) ≤ h*(n),     h(p) ≤ c(p,q) + h(q)')
    p('h=0 时 A* 退化为一致代价搜索。独立 Dijkstra 基准使用未确定节点集合和线性最小值扫描，没有调用 A* 或复用其优先队列。随机小图另与 BFS/Bellman-Ford 对照，测试逐格检查启发下界和邻边一致性。起终点相同返回单点路径或费用 0；障碍起终点视为不可达，负权和非法地图明确拒绝。')
    heading('3.3 时间与空间开销',2)
    p('四邻网格满足 E=O(V)。二叉堆实现通常为 O((V+E)log V)，距离及父节点记录为 O(V)；重复入堆会增加堆项数，故同时记录堆操作与过期出队。结果中的“扩展节点数”只计搜索展开，不包含启发预处理访问，预处理工作量与耗时另列。')

    page();heading('4 题目一 迷宫最短路径')
    path=examples['q1_manhattan']['output']
    p('官方地图起点为 (0,0)，终点为 (3,4)。父节点链回溯得到 '+ '→'.join(f'({x},{y})' for x,y in path)+'，共 7 步。起终点曼哈顿距离为 7，因此该合法路径达到步数下界，已经是最短路径。')
    table(['启发','生成','扩展','有效出队','堆调用','核心 / IQR μs','总时间 μs'],[
        [LABEL[m],r('maze_official',m)['generated'],r('maze_official',m)['expanded'],r('maze_official',m)['visited'],
         r('maze_official',m)['heap_operations'],interval(r('maze_official',m)),us(r('maze_official',m),'total_ns')]
        for m in ('zero','manhattan','topology')], [2.4,1.3,1.3,1.7,1.6,4.2,4.3], '表5 官方迷宫搜索开销',fontsize=8.7)
    p('单位代价下，曼哈顿距离是忽略障碍后的步数下界，阻止搜索过早探索远离目标的候选。h=0 展开 15 个节点，曼哈顿展开 13 个；目标计入出队而不计入扩展，故对应出队数为 16 和 14。拓扑下界保留障碍，进一步降至 7 次扩展，但要先从目标执行一次 BFS。')
    figure('astar_costs.png','图1 A* 启发比较与障碍密度敏感性',16.2)
    p(f"拓扑模式的核心中位时间为 {us(r('maze_official','topology'))} μs，计入预处理后为 {us(r('maze_official','topology'),'total_ns')} μs；曼哈顿总中位时间为 {us(r('maze_official','manhattan'),'total_ns')} μs。该实例说明减少扩展并不自动转化为单次端到端加速。图1下半部分包含不可达样本，应结合第10节的可达数量解释。")

    page();heading('5 题目二 加权地形与可采纳启发')
    heading('5.1 进入费用与官方最优路径',2)
    p('官方起点为 (0,0)，终点为 (4,4)。每次移动收取目标格权值，起点不计费，进入 E 的费用为 0。S 内部按 1 规范化只影响再次进入起点；非负代价下存在不重访起点的最优简单路径，因此不改变 S 到 E 的最优结果。')
    weighted_path=json.loads(next(row['path_json'] for row in __import__('audit_results').read_csv(ROOT/'results/benchmarks/astar.csv') if row['case']=='weighted_official' and row['mode']=='manhattan'))
    cumulative=0; cost_rows=[]
    for index,(row,col) in enumerate(weighted_path):
        step=0 if index==0 or TERRAIN[row][col]=='E' else TERRAIN[row][col]
        cumulative+=step
        cost_rows.append([index,f'({row},{col})',step,cumulative])
    table(['步次','位置','本步进入费用','累计 g'],cost_rows,[2.3,4.5,5,5], '表6 官方加权路径的费用累计',fontsize=9)
    heading('5.2 为什么不能直接使用曼哈顿距离',2)
    p('最小反例为一行 [S,E]：真实最优费用为 0，而曼哈顿距离为 1，发生高估。设 d(n,E) 为曼哈顿距离，m 为非目标可通行格的最小进入费用，cE 为目标进入费用。采用如下修正，目标处单独设为 0。')
    equation('h_M(E)=0;    h_M(n)=m·max(d(n,E)−1,0)+cE  (n≠E)')
    p('任意从非目标 n 到 E 的路径至少走 d 步，其中最后一步费用是 cE，其余至少 d−1 次非目标进入各不小于 m，故 hM 不高估。非目标邻格间 d 最多改变 1，h 的下降至多为 m，不大于进入费用；若邻格为 E，则 hM(n)=cE。因此启发函数可采纳且一致。若存在零权非目标格，则 m=0，在 E=0 时退化为 h=0。')
    table(['启发','最优费用','扩展','有效出队','核心 / IQR μs','总时间 μs'],[
        [LABEL[m],r('weighted_official',m)['cost'],r('weighted_official',m)['expanded'],r('weighted_official',m)['visited'],interval(r('weighted_official',m)),us(r('weighted_official',m),'total_ns')]
        for m in ('zero','manhattan','topology')], [2.5,2,2,2.1,4.4,3.8], '表7 官方加权地形结果',fontsize=8.7)
    p('该实例 hM(S)=7，总代价也为 7。拓扑模式与修正曼哈顿均展开 10 个节点，未进一步减少扩展，却需要额外 BFS。加权图的“最少步数”不等于“最小费用”；两种启发都只作下界，最终费用仍由真实进入权值累加。')

    page();heading('6 Minimax与Alpha-Beta原理')
    p('TreeNode 使用 children 表示有序子节点，叶节点用 value 保存分数；非叶值由递归计算。根为 MAX，之后 MAX/MIN 交替；终止叶直接返回给定分数。完整 Minimax 遍历所有叶，值回传遵循：')
    equation('V(n)=value(n)  (leaf);   V(n)=max V(child)  (MAX);   V(n)=min V(child)  (MIN)')
    p('Alpha-Beta 在相同递归上维护 α 与 β。α 是 MAX 在祖先分支中已能保证的下界，β 是 MIN 已能保证的上界。MAX 用已搜索子值增大 α，MIN 用已搜索子值减小 β；当 α≥β 时，其余兄弟无法改善祖先决策，可以跳过。顺序改变可能更早产生紧边界，但不会改变完整根值。')
    p('实现只有在确实存在尚未访问兄弟时增加剪枝事件数。被跳过的直接分支数与其完整子树节点数分别统计；不把一次剪枝误记为一个未访问节点。没有剪枝时，时间为 O(b^H)，递归栈深度为 O(H)；理想排序可显著缩减访问，但最坏情况仍可能接近完整搜索。')
    heading('7 题目三 官方博弈树与顺序实验')
    heading('7.1 数值回传和一次真实剪枝',2)
    p('按原文第10—12页构造，左 MIN 的子树为 MAX(3,5) 和 MAX(6,9)，右 MIN 的子树为 MAX(6,9) 和 MAX(1,2)。第12页 max2 被共享引用，但搜索按每条路径展开，不按对象身份去重。')
    equation('max( min(max(3,5),max(6,9)), min(max(6,9),max(1,2)) ) = max(5,2) = 5')
    p('从左到右时，左 MIN 先得到 5，从而 β=5。随后进入 MAX(6,9)，首先读到 6，已满足 α=6≥β=5，于是无需读取其右叶 9。返回的 6 是该分支的下界，不是完整子树精确值 9，但足以证明左 MIN 保持 5。此处不能把剪枝返回值写成完整子树的真实最优值。')
    table(['模式','根值','访问','搜索叶','剪枝事件','跳过节点','排序叶读取'],[
        [LABEL[m],r('tree_official',m)['value'],r('tree_official',m)['visited'],r('tree_official',m)['leaf_evaluations'],r('tree_official',m)['cutoffs'],r('tree_official',m)['skipped_nodes'],r('tree_official',m)['ordering_leaves']]
        for m in ('minimax','left','right','ordered')],[3.5,1.4,1.5,1.8,2.6,2.4,3.6],'表8 官方树的值与工作量',fontsize=9)
    p('原始 Alpha-Beta 接口严格从左到右，返回 5，访问 14 次、叶评价 7 次；完整 Minimax 分别为 15 和 8。逆序与样本排序仅作为实验扩展，不改变题目三默认遍历约定。')

    page();heading('7.2 访问顺序的收益与额外成本',2)
    p('排序方案为每个候选子树读取最左叶作为样本，MAX 优先较大样本，MIN 优先较小样本。该信息不是完整 Minimax 解，因而可能误导排序；读取样本需要沿子树下降，计入 ordering_nodes、ordering_leaves 和核心耗时。')
    figure('tree_orders.png','图2 访问顺序 深度与额外样本叶成本',16.2)
    table(['深度7平衡树','访问','搜索叶','排序叶','核心 μs','含结构统计 μs'],[
        [LABEL[m],r('tree_balanced_d7',m)['visited'],r('tree_balanced_d7',m)['leaf_evaluations'],r('tree_balanced_d7',m)['ordering_leaves'],us(r('tree_balanced_d7',m)),us(r('tree_balanced_d7',m),'total_ns')]
        for m in ('minimax','left','right','ordered')],[3.6,2,2,2,3.5,3.7],'表9 深度7平衡树的排序代价',fontsize=9)
    p('样本排序将访问数从 1,476 降至 875，但额外读取 1,284 次样本叶，正式叶评价 447 次，合计 1,731 次叶读取，超过原顺序的 855 次。逆序访问仅为 628，说明该样本与真实子树极值并不总是吻合。六棵额外树中排序均比左到右减少访问，但不能由这六个实例推出普遍更快。')
    p('结构大小遍历只用于准确计算被跳过节点数，对四种模式统一执行并单列预处理时间。报告同时呈现核心时间与含统计预处理的时间，避免将测量工具自身的工作误判为搜索算法的必需成本。')

    page();heading('8 题目四 地图 距离矩阵与状态模型')
    table(['r/c']+list(range(7)),[[i]+row for i,row in enumerate(TREASURE_MAP)],
          [2.1]*8,'表10 原题7×7地图',fontsize=9)
    p('A=(0,0)，B=(6,6)；T1=(0,6)、价值9，T2=(2,2)、价值6，T3=(4,4)、价值7，T4=(6,0)、价值7。A/B/T 所在格的进入费用均为1，与题二 E=0 的规则不同。其他玩家、已获取和未获取宝藏的格子始终可通行，经过宝藏不自动收集。')
    labels,matrix=examples['q4_distances']['output']
    table(['起点/终点']+labels,[[label]+row for label,row in zip(labels,matrix)],
          [2.4]*7,'表11 A* 构造的6×6最小移动费用矩阵',fontsize=9)
    p('36 个有序点对均独立调用 A*，并通过独立最短路对照。对角线为0。虽然进入格代价在一般加权图上具有方向性，本题六个端点的格费都为1；同一路径正反向费用之差仅为端点进入费之差，因此本矩阵对称，不能据此假设任意加权图距离都对称。')
    equation('s=(p_A,p_B,R,Score_A,Score_B,turn)')
    equation('Gain(p,T_i)=V_i−d(p,T_i),     U(s)=Score_A−Score_B')
    p('GameState 为不可变状态。每步从 R 中只移除目标宝藏，更新当前行动者的位置与累计收益，另一方位置保持不变，再切换轮次。A 最大化 U，B 最小化 U；每层是一名玩家的一次行动，达到 H 或 R 为空即评价当前 U，不追加未来收益猜测。固定矩阵可以复用，是因为共享地图的通行费用不随占位和拾取改变。')

    page();heading('8.2 搜索深度与完整决策分析',2)
    branch_rows=[]
    for row in branches:
        seq=' → '.join(f"{x['player']}:{x['target']}({x['gain']:+d})" for x in row['line'])
        branch_rows.append([row['depth'],row['first'],row['value'],seq])
    table(['H','A首步','该首步精确值','一条双方最优应对序列  括号内为Gain'],branch_rows,
          [1.1,1.7,2.4,11.6],'表12 固定首步后完整Minimax求得的精确分支值',fontsize=9)
    p('表12直接取自 decision_branches.json，并重新按矩阵逐步计算累计分数。先固定首步，再以完整 Minimax 求余下深度，避免将窄窗口剪枝边界当作分支精确值。H=2 的四个首步值为 −4、−1、−2、−4，故选 T2；H=4 为 −5、1、3、−7，故选 T3。')
    p('H=2 时，A 到 T2 费用4、获得6，净收益2；B 随后从 B 到 T3 费用4、获得7，净收益3。因此 U=2−3=−1。若 A 首取 T3，则收益为−1，B 取 T1 得1，当前差值为−2，短视评价不支持首取 T3。')
    best4=next(x for x in branches if x['depth']==4 and x['first']=='T3')
    table(['轮次','玩家与移动','宝藏价值','移动费','Gain','A累计','B累计','U'],[
        [i+1,x['player']+':'+x['source']+'→'+x['target'],x['treasure_value'],x['cost'],x['gain'],x['score_a'],x['score_b'],x['utility']]
        for i,x in enumerate(best4['line'])],[1.2,4.2,2,1.8,1.7,1.9,1.9,2.1],'表13 H=4 首取T3的一条最优收益展开',fontsize=9)
    p('完整四回合中，A 虽首轮损失1，但可在第二次行动后使 B 从 T1 前往 T4，移动费12超过价值7，B 最后一轮 Gain=−5。A 总收益为−1，B总收益为−4，故 U=3。首取T2的完整分支则为 A：2+3=5，B：3+1=4，U=1。因此选择变化来自后续位置与宝藏分配，而不是修改评价函数。')
    p('B 的后续选择同样经过最小化：这里只展示一条稳定顺序选出的最优应对，不声称它是唯一序列。当前行动即时 Gain 排序只是搜索顺序提示，不能替代完整四回合决策。')

    page();heading('9 排序 缓存与组合优化对照')
    modes=('minimax','left','right','ordered','cache','ordered_cache')
    table(['方法','H2访问/叶','H4访问/叶','H4剪枝','H4核心/IQR μs','H4总 μs'],[
        [LABEL[m],f"{r('treasure_H2',m)['visited']}/{r('treasure_H2',m)['leaf_evaluations']}",
         f"{r('treasure_H4',m)['visited']}/{r('treasure_H4',m)['leaf_evaluations']}",r('treasure_H4',m)['cutoffs'],interval(r('treasure_H4',m)),us(r('treasure_H4',m),'total_ns')]
        for m in modes],[3.5,2.4,2.4,1.7,4,2.8],'表14 官方四宝藏算法消融',fontsize=8.8)
    p('所有方法在两个深度均保持相同的 best_target 和 best_value。即时收益排序使 H=4 访问从47降至37，减少21.3%；叶评价从16降至12。逆序H=4发生一次平局候选的全窗口重搜，因此访问数60与跳过节点21不能直接相加当成原始65节点树的互补划分。')
    figure('treasure_costs.png','图3 官方综合题的访问与叶评价对照',15.8)
    p('H=2 的缓存原顺序为0/12次命中，排序缓存为0/10；H=4分别为0/46、0/36，命中率均为0%。完整状态键较保守且四回合路径短，本实例未产生可复用的相同状态。缓存保持37或47次访问，却增加哈希、查表和存储开销，不能宣称获得加速。')

    page();heading('10 创新拓展与参数敏感性')
    heading('10.1 可证明的拓扑启发与消融',2)
    p('拓扑启发先以 BFS 得到保留障碍后的最少步数 L，再对所有非目标步统一使用最小进入费 m：')
    equation('h_T(E)=0;    h_T(n)=m·(L(n,E)−1)+cE')
    p('任意真实路径不少于 L 步，因此该式为费用下界；邻格 L 相差至多1，从而满足一致性。不可达分量使用0。在单位迷宫中，BFS已计算精确步数，单次查询必须承担这份预处理成本；它可能适合固定终点多次查询，但本实验未测量摊销收益。零权使m=0时也可能失去区分能力。')
    figure('ablation.png','图4 启发 排序 缓存及组合的消融实验',16.2)
    p('折返迷宫中三种启发均扩展285次，额外BFS只增加成本。该长通道与FIFO同值顺序形成搜索平台，启发更强仍可能访问同样多的节点。综合题的组合优化比单独排序多了缓存维护而没有命中收益，因此组合不必优于单项。上述方法属于已有搜索思想在本题的实现与实证比较，不主张算法原创性。')

    page();heading('10.2 障碍密度 权重和树深度的影响',2)
    sensitivity=[]
    for kind in ('maze','weighted'):
        for density in (.1,.25,.4):
            for weight in ((1,) if kind=='maze' else (3,9)):
                group=[x for x in data.values() if x['family']==kind and x.get('density')==density and x.get('max_weight')==weight]
                count=sum(x['cost']>=0 for x in group if x['mode']=='zero')
                means=[f"{statistics.mean(x['expanded'] for x in group if x['mode']==m):.2f}" for m in ('zero','manhattan','topology')]
                sensitivity.append(['迷宫' if kind=='maze' else '加权',density,weight,f'{count}/3']+means)
    table(['类型','障碍概率','权重上限','可达数','h=0','曼哈顿','拓扑'],sensitivity,
          [2,2.4,2.4,2.2,2.6,2.6,2.6],'表15 三个种子的平均扩展数 包含不可达图',fontsize=9)
    p('障碍概率0.40的三个种子均不可达，平均扩展降至75.33是因为起点连通分量缩小，而不是高密度障碍普遍更易求解。概率0.25仅一张可达，平均值同样混合了不同问题状态；不能据密度曲线单独推断可达问题的难度。')
    p('低密度加权图中，权重上限3时，曼哈顿和拓扑平均扩展分别为429.67、410.67；上限9时为489.67、485.67。最小步费仍为1，真实费用增大后下界相对变弱，减少的扩展趋少。这里只能说明这些种子上的现象，未作统计总体推断。')
    table(['树类型','深度','完整访问','原顺序AB','逆序AB','样本排序AB'],[
        ['平衡' if kind=='balanced' else '非平衡',depth]+[r(f'tree_{kind}_d{depth}',m)['visited'] for m in ('minimax','left','right','ordered')]
        for kind in ('balanced','unbalanced') for depth in (3,5,7)],
        [3,1.5,3.2,3.1,3,3],'表16 自设树的深度对照',fontsize=9)
    p('平衡树完整访问满足 (3^(H+1)−1)/2，H=3/5/7 分别为40、364、3280，与记录一致。剪枝比例受叶值分布及访问顺序影响。不同深度使用不同固定种子，非平衡树也可能提前终止，因此本组不是同一棵树的逐层截断，不能将所有差异都归因于深度。')

    page();heading('11 算法行为与正确性分析')
    heading('11.1 缓存上下界与稳定平局',2)
    p('置换表键为完整状态和剩余深度，包含双方位置、剩余宝藏、双方累计分数及轮次。叶值为EXACT；内部返回值v按原调用窗口分类：v≤α0记UPPER_BOUND，v≥β0记LOWER_BOUND，窗口内记EXACT。读取上下界只收紧窗口；只有窗口关闭才直接返回边界，不将截断值当作精确最优值。')
    p('同值决策按原remaining顺序稳定选取，官方即T1、T2、T3、T4。当窄窗口结果与当前最好值相等且该候选排序更早时，先用全窗口重搜，再决定是否替换。否则一个仅为上界的“等值”可能错误改变best_target。窗口复用、完整状态分数偏移、不同深度及随机平局博弈均有回归测试。')
    heading('11.2 证据链与审查结果',2)
    p(f"最终完整测试为 {tests.group(1)} passed。62项测试包括原有55项和新增7项输入/统计回归测试；四题官方结果、全部扩展模式和2220条性能记录均通过最优性校验。审查脚本独立从原始CSV重新计算中位数、线性插值四分位数、IQR与计数，核对148组汇总、58行分析表、9行敏感性表和8条首步分支。")
    p('发现并修复的问题是实验扩展接口没有完整拒绝非法深度、轮次和关闭的α/β窗口，以及原汇总只检查部分字段稳定。修复后将所有非计时输出与计数纳入重复一致性检查，重新运行测试和实验。官方语义及最优值不变，前一轮原始记录保留为历史数据，不与本轮耗时混用。')
    heading('12 实验局限与总结')
    p('样本规模有限：每个地图参数组仅三个种子，每种树形深度仅一个实例；没有控制频率、进程调度或GC，也没有作显著性检验。排序内部计时本身存在开销，树结构遍历属于统计辅助，解释微秒差异时需谨慎。缓存正确性在更多状态复用测试中验证，但官方实例零命中，不能据此证明其在更大问题上一定有效或无效。')
    p('实验确认：A*应从真实移动费用设计下界，E=0使原始曼哈顿高估；Alpha-Beta的核心收益取决于边界形成顺序，样本排序的额外成本可能抵消节点下降；四宝藏决策必须同时考虑双方位置和未来收益，H=4的选择不同于H=2。性能结论应联合查看工作量、预处理、计时分布与正确性，而非仅比较平均速度。')
    heading('资料与复现入口',2)
    p('任务依据：《当代人工智能实验课 实验二 搜索算法》，2026年10月15日，第1—17页。实验依据：results/benchmarks/ 下三份原始CSV、config.json、decision_branches.json；图1—4由scripts/plot_results.py生成。')
    p('复现顺序：安装requirements.txt；python -m pytest -q；python scripts/run_benchmarks.py；python scripts/plot_results.py；python scripts/audit_results.py。报告构建使用requirements-report.txt中的python-docx，PDF由Word导出。正式提交目录附有源码、测试、脚本、报告与原始数据。')

    folder=ROOT/'report';folder.mkdir(exist_ok=True)
    output=folder/'实验二_搜索算法_实验报告.docx'
    doc.save(output)
    print(f'Created {output}; independent data audit: {audit_result["all_checks_passed"]}')


if __name__=='__main__':
    main()
